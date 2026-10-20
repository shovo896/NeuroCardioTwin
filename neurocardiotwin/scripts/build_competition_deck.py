"""Create the NeuroCardioTwin competition presentation.

Run from the repository root:
  .venv/bin/python scripts/build_competition_deck.py
"""
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


OUT = Path(__file__).resolve().parents[1] / "NeuroCardioTwin_Competition_Presentation.pptx"
W, H = Inches(13.333), Inches(7.5)
NAVY = RGBColor(7, 17, 34)
PANEL = RGBColor(18, 39, 64)
PANEL_2 = RGBColor(22, 48, 76)
CYAN = RGBColor(34, 211, 238)
TEAL = RGBColor(20, 184, 166)
WHITE = RGBColor(237, 246, 255)
MUTED = RGBColor(164, 190, 214)
GREEN = RGBColor(92, 225, 167)
AMBER = RGBColor(251, 191, 36)
VIOLET = RGBColor(167, 139, 250)
RED = RGBColor(251, 113, 133)


def shape(slide, kind, x, y, w, h, fill, radius=True, line=None):
    s = slide.shapes.add_shape(kind if not radius else MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = fill
    s.line.color.rgb = line or fill
    return s


def text(slide, value, x, y, w, h, size=18, color=WHITE, bold=False, font="Aptos", align=None, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame; frame.clear(); frame.word_wrap = True; frame.vertical_anchor = valign
    p = frame.paragraphs[0]; p.text = value; p.font.name = font; p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = color
    if align is not None: p.alignment = align
    return box


def title(slide, kicker, heading, sub=""):
    text(slide, kicker.upper(), Inches(.62), Inches(.40), Inches(11.8), Inches(.26), 9, CYAN, True, "Courier New")
    text(slide, heading, Inches(.60), Inches(.72), Inches(12.0), Inches(.58), 27, WHITE, True)
    if sub: text(slide, sub, Inches(.62), Inches(1.34), Inches(11.6), Inches(.48), 12, MUTED)
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.62), Inches(1.93), Inches(12.1), Inches(.02)).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = RGBColor(48, 84, 119); slide.shapes[-1].line.fill.background()


def card(slide, heading, body, x, y, w, h, accent=CYAN, icon=""):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, PANEL, True, RGBColor(47, 84, 118))
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(.20), y + Inches(.20), Inches(.48), Inches(.48), accent)
    text(slide, icon, x + Inches(.20), y + Inches(.27), Inches(.48), Inches(.22), 15, NAVY, True, align=PP_ALIGN.CENTER)
    text(slide, heading, x + Inches(.82), y + Inches(.22), w - Inches(1.0), Inches(.28), 14, WHITE, True)
    text(slide, body, x + Inches(.22), y + Inches(.84), w - Inches(.44), h - Inches(1.0), 11.5, MUTED)


def bullets(slide, entries, x, y, w, font_size=15, gap=.48):
    for index, entry in enumerate(entries):
        yy = y + Inches(index * gap)
        text(slide, "●", x, yy, Inches(.22), Inches(.24), 10, CYAN)
        text(slide, entry, x + Inches(.28), yy - Inches(.015), w - Inches(.28), Inches(.38), font_size, WHITE)


def footer(slide, page):
    text(slide, "NEUROCARDIOTWIN  •  AI-POWERED BRAIN–HEART DIGITAL TWIN", Inches(.62), Inches(7.13), Inches(8), Inches(.16), 7, MUTED, True, "Courier New")
    text(slide, f"{page:02d}", Inches(12.15), Inches(7.09), Inches(.55), Inches(.2), 9, CYAN, True, "Courier New", PP_ALIGN.RIGHT)


def make_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); slide.background.fill.solid(); slide.background.fill.fore_color.rgb = NAVY
    return slide


def main():
    prs = Presentation(); prs.slide_width = W; prs.slide_height = H

    # 1 — Title
    s = make_slide(prs)
    shape(s, MSO_SHAPE.OVAL, Inches(9.4), Inches(-1.2), Inches(5.0), Inches(5.0), RGBColor(15, 62, 88), False)
    shape(s, MSO_SHAPE.OVAL, Inches(10.0), Inches(-.6), Inches(3.9), Inches(3.9), RGBColor(12, 86, 104), False)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.67), Inches(.75), Inches(.74), Inches(.74), TEAL)
    text(s, "⌁", Inches(.67), Inches(.83), Inches(.74), Inches(.35), 31, NAVY, True, align=PP_ALIGN.CENTER)
    text(s, "NEUROCARDIOTWIN", Inches(.68), Inches(1.82), Inches(8), Inches(.38), 13, CYAN, True, "Courier New")
    text(s, "Signals, with\ncontext.", Inches(.65), Inches(2.25), Inches(8), Inches(1.45), 40, WHITE, True)
    text(s, "A privacy-first AI brain–heart digital twin for\nresearch-grade EEG and ECG decision support.", Inches(.70), Inches(3.88), Inches(6.8), Inches(.7), 17, MUTED)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.70), Inches(5.15), Inches(3.35), Inches(.5), PANEL_2, True, CYAN)
    text(s, "বাংলা + English • Secure • Explainable", Inches(.88), Inches(5.30), Inches(3.0), Inches(.2), 10, GREEN, True)
    text(s, "Competition Pitch Deck  |  2026", Inches(.70), Inches(6.55), Inches(5), Inches(.25), 10, MUTED)
    footer(s, 1)

    # 2 — Problem
    s = make_slide(prs); title(s, "The problem", "Vital signals arrive disconnected.", "Clinicians and researchers need a safer bridge from raw biosignal to actionable research workflow.")
    card(s, "Fragmented modalities", "EEG and ECG are commonly reviewed in separate tools, obscuring brain–heart context.", Inches(.65), Inches(2.35), Inches(3.8), Inches(2.25), CYAN, "↔")
    card(s, "Opaque AI outputs", "A probability alone cannot tell a researcher what to verify next or where the evidence came from.", Inches(4.77), Inches(2.35), Inches(3.8), Inches(2.25), AMBER, "?")
    card(s, "Trust & access gap", "Research platforms must protect private data while remaining usable for Bangla-speaking users.", Inches(8.9), Inches(2.35), Inches(3.8), Inches(2.25), VIOLET, "✓")
    text(s, "Design question", Inches(.68), Inches(5.25), Inches(2), Inches(.25), 10, CYAN, True, "Courier New")
    text(s, "How do we make multimodal AI understandable, privacy-aware, and inclusive—without pretending it is a medical device?", Inches(.68), Inches(5.58), Inches(11.5), Inches(.65), 22, WHITE, True)
    footer(s, 2)

    # 3 — Solution
    s = make_slide(prs); title(s, "Our solution", "One secure workspace. Three intelligent layers.", "NeuroCardioTwin turns validated signal inputs into a grounded research decision—not a diagnosis.")
    for idx, (head, body, color, num) in enumerate([
        ("Local inference", "ONNX EEG seizure, ECG rhythm, and joint forecasting models run in the application backend.", CYAN, "01"),
        ("Evidence grounding", "Hugging Face embeddings + Pinecone retrieve only approved research knowledge before AI guidance is generated.", TEAL, "02"),
        ("Human-centred delivery", "Bilingual UI, clear safety boundaries, private profile/history, and a live research assistant.", VIOLET, "03"),
    ]):
        x = Inches(.72 + idx * 4.15)
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.35), Inches(3.65), Inches(3.2), PANEL, True, RGBColor(46, 84, 117))
        text(s, num, x + Inches(.25), Inches(2.60), Inches(.65), Inches(.32), 13, color, True, "Courier New")
        text(s, head, x + Inches(.25), Inches(3.22), Inches(3.05), Inches(.35), 17, WHITE, True)
        text(s, body, x + Inches(.25), Inches(3.78), Inches(3.05), Inches(1.0), 12, MUTED)
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(.25), Inches(4.95), Inches(2.8), Inches(.05), color)
    footer(s, 3)

    # 4 — Architecture
    s = make_slide(prs); title(s, "System architecture", "A privacy-first decision pipeline.", "Raw waveforms stay inside the controlled model path; only compact output summaries enter the RAG workflow.")
    stages = [("1", "Validated input", "EEG / ECG\nCSV or sensor demo", CYAN), ("2", "Local ONNX", "Inference +\ninput-quality checks", TEAL), ("3", "Private RAG", "HF embedding →\nPinecone retrieval", VIOLET), ("4", "Decision support", "Grounded response +\nsource citations", AMBER)]
    for i, (n, h, b, c) in enumerate(stages):
        x = Inches(.65 + i * 3.12)
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.7), Inches(2.5), Inches(1.75), PANEL, True, c)
        text(s, n, x + Inches(.18), Inches(2.92), Inches(.35), Inches(.2), 10, c, True, "Courier New")
        text(s, h, x + Inches(.20), Inches(3.32), Inches(2.1), Inches(.25), 14, WHITE, True)
        text(s, b, x + Inches(.20), Inches(3.75), Inches(2.05), Inches(.52), 10.5, MUTED)
        if i < 3: text(s, "→", x + Inches(2.62), Inches(3.36), Inches(.38), Inches(.25), 21, CYAN, True, align=PP_ALIGN.CENTER)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.46), Inches(5.32), Inches(10.4), Inches(.65), RGBColor(14, 55, 65), True, RGBColor(37, 133, 139))
    text(s, "Privacy boundary: no email, access token, profile data, upload, or raw waveform is sent to the RAG prompt or vector index.", Inches(1.72), Inches(5.54), Inches(9.9), Inches(.2), 11, GREEN, True, align=PP_ALIGN.CENTER)
    footer(s, 4)

    # 5 — Live workflow
    s = make_slide(prs); title(s, "User journey", "From signal upload to a transparent research output.")
    bullets(s, ["Sign in securely with Auth0; each user has a protected private workspace.", "Upload numeric EEG/ECG data or select a controlled demo signal.", "Run local inference with explicit sample-validation and model confidence.", "Receive an integrated decision with approved evidence sources and safety boundaries.", "Review only your own compact decision history—never raw waveform history."], Inches(.85), Inches(2.28), Inches(7.0), 14, .63)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.55), Inches(2.30), Inches(3.75), Inches(3.4), PANEL, True, RGBColor(49, 87, 121))
    text(s, "LIVE DEMO", Inches(8.82), Inches(2.60), Inches(2), Inches(.2), 10, CYAN, True, "Courier New")
    text(s, "1 minute", Inches(8.82), Inches(3.05), Inches(2.5), Inches(.4), 29, WHITE, True)
    text(s, "Sign in → upload/demo → run model → see evidence-grounded decision → ask the assistant", Inches(8.84), Inches(3.75), Inches(2.9), Inches(.95), 13, MUTED)
    text(s, "Judges can inspect the result, source citations, and privacy story in one flow.", Inches(8.84), Inches(5.05), Inches(2.9), Inches(.38), 10.5, GREEN, True)
    footer(s, 5)

    # 6 — Differentiation
    s = make_slide(prs); title(s, "Why this is different", "Not just another prediction dashboard.")
    entries = [("Multimodal by design", "EEG, ECG, and their future trajectory live in one research workspace.", CYAN), ("RAG is integrated", "Evidence retrieval shapes the same decision workflow—not a disconnected chatbot.", TEAL), ("Safe by construction", "Prototype framing, source-aware responses, and no diagnostic claims.", AMBER), ("Bangla-ready access", "The full static interface and assistant can be switched between English and বাংলা.", VIOLET)]
    for i, (h, b, c) in enumerate(entries):
        x = Inches(.7 + (i % 2) * 6.05); y = Inches(2.3 + (i // 2) * 1.75)
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.55), Inches(1.35), PANEL, True, RGBColor(49, 85, 118))
        shape(s, MSO_SHAPE.OVAL, x + Inches(.25), y + Inches(.33), Inches(.55), Inches(.55), c, False)
        text(s, h, x + Inches(1.0), y + Inches(.25), Inches(4.2), Inches(.24), 15, WHITE, True)
        text(s, b, x + Inches(1.0), y + Inches(.66), Inches(4.15), Inches(.38), 11, MUTED)
    footer(s, 6)

    # 7 — Trust
    s = make_slide(prs); title(s, "Trust, safety & privacy", "The strongest AI system is the one that knows its boundary.")
    bullets(s, ["Auth0-protected application access and scoped MCP permissions.", "Server-side secret handling: browser never receives Hugging Face, Pinecone, or Supabase service keys.", "Per-user, server-filtered history with Row-Level Security; compact metadata only.", "RAG index contains reviewed local knowledge—not personal records or waveforms.", "Visible research-only safety notice; no diagnosis, triage, or treatment instruction."], Inches(.85), Inches(2.25), Inches(7.1), 14, .62)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.55), Inches(2.30), Inches(3.7), Inches(3.25), RGBColor(21, 61, 58), True, RGBColor(52, 148, 129))
    text(s, "✓", Inches(9.72), Inches(2.75), Inches(1.3), Inches(.75), 42, GREEN, True, align=PP_ALIGN.CENTER)
    text(s, "Privacy is a feature,\nnot a footnote.", Inches(8.92), Inches(3.85), Inches(3.0), Inches(.65), 20, WHITE, True, align=PP_ALIGN.CENTER)
    text(s, "Designed so researchers can explain what crosses the AI boundary—and what does not.", Inches(8.95), Inches(4.78), Inches(2.95), Inches(.43), 10.5, MUTED, False, align=PP_ALIGN.CENTER)
    footer(s, 7)

    # 8 — Safety escalation
    s = make_slide(prs); title(s, "Safety escalation", "AI is not the last stop in an emergency.", "NeuroCardioTwin keeps the emergency path separate from model inference and assistant chat.")
    card(s, "Recognise", "Visible urgent-warning panel covers chest pain or pressure, severe breathing difficulty, fainting, confusion, seizure and severe bleeding.", Inches(.72), Inches(2.35), Inches(3.72), Inches(2.55), AMBER, "!")
    card(s, "Act", "One tap opens the phone dialler for Bangladesh National Emergency Service 999. The interface tells users not to wait for an AI reply.", Inches(4.82), Inches(2.35), Inches(3.72), Inches(2.55), RGBColor(225, 29, 72), "☎")
    card(s, "Support", "Bilingual Bengali/English instructions: alert someone nearby, do not stay alone, keep the phone available, and avoid driving when severely unwell.", Inches(8.92), Inches(2.35), Inches(3.72), Inches(2.55), VIOLET, "✓")
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.22), Inches(5.45), Inches(10.1), Inches(.6), RGBColor(63, 24, 37), True, RGBColor(180, 66, 87))
    text(s, "Safety boundary: the app never diagnoses an emergency; it makes the fastest real-world help path obvious.", Inches(1.5), Inches(5.66), Inches(9.55), Inches(.22), 12, RGBColor(255, 218, 224), True, align=PP_ALIGN.CENTER)
    footer(s, 8)

    # 9 — Validation
    s = make_slide(prs); title(s, "What we validate", "Real model execution—not invented metrics.", "Model performance is reported from the connected backend, while the interface exposes input and runtime constraints.")
    card(s, "Input integrity", "Required sample windows, finite-value checks, sampling rate, and validated input state before inference.", Inches(.7), Inches(2.35), Inches(3.8), Inches(2.2), CYAN, "✓")
    card(s, "Operational proof", "ONNX CPU model status, output confidence, classification/forecast details, and API-backed metrics surfaces.", Inches(4.78), Inches(2.35), Inches(3.8), Inches(2.2), TEAL, "↗")
    card(s, "Evidence traceability", "Decision outputs cite approved knowledge titles, making the RAG layer inspectable by judges.", Inches(8.86), Inches(2.35), Inches(3.8), Inches(2.2), VIOLET, "#")
    text(s, "Honest claim", Inches(.72), Inches(5.23), Inches(1.8), Inches(.22), 10, AMBER, True, "Courier New")
    text(s, "This is a research prototype, not a clinical device. Its strength is a transparent, reproducible workflow—not overstated medical claims.", Inches(.72), Inches(5.55), Inches(11.1), Inches(.48), 18, WHITE, True)
    footer(s, 9)

    # 10 — Impact & roadmap
    s = make_slide(prs); title(s, "Impact & next steps", "A platform designed to grow with the research community.")
    bullets(s, ["Pilot with supervised research teams using consented, de-identified signal datasets.", "Expand the approved knowledge base through clinician/researcher review workflows.", "Add calibrated model evaluation, cohort-level fairness analysis, and reproducible audit exports.", "Connect validated real-time devices only after robust signal-quality and governance testing."], Inches(.85), Inches(2.25), Inches(6.9), 14, .68)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.25), Inches(2.25), Inches(4.0), Inches(3.4), PANEL, True, RGBColor(49, 87, 121))
    text(s, "VISION", Inches(8.55), Inches(2.57), Inches(2), Inches(.2), 10, CYAN, True, "Courier New")
    text(s, "Inclusive,\nevidence-aware\nbrain–heart AI.", Inches(8.55), Inches(3.06), Inches(3.0), Inches(1.25), 24, WHITE, True)
    text(s, "Built in Bangladesh; designed for globally responsible digital-health research.", Inches(8.55), Inches(4.86), Inches(3.0), Inches(.48), 11, GREEN, True)
    footer(s, 10)

    # 11 — Closing
    s = make_slide(prs)
    shape(s, MSO_SHAPE.OVAL, Inches(-1.3), Inches(3.9), Inches(4.1), Inches(4.1), RGBColor(14, 78, 94), False)
    text(s, "NEUROCARDIOTWIN", Inches(.7), Inches(1.15), Inches(6), Inches(.25), 11, CYAN, True, "Courier New")
    text(s, "Better signals.\nBetter context.\nBetter research decisions.", Inches(.65), Inches(1.65), Inches(9.6), Inches(1.75), 37, WHITE, True)
    text(s, "The brain and heart should not be understood in isolation.", Inches(.70), Inches(3.78), Inches(7.7), Inches(.35), 16, MUTED)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.72), Inches(4.75), Inches(3.45), Inches(.62), TEAL)
    text(s, "Thank you  •  Questions?", Inches(.98), Inches(4.96), Inches(2.95), Inches(.23), 13, NAVY, True, align=PP_ALIGN.CENTER)
    text(s, "Live demo: neurocardiotwin.vercel.app", Inches(.72), Inches(6.25), Inches(5.4), Inches(.25), 11, CYAN, True, "Courier New")
    footer(s, 11)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
